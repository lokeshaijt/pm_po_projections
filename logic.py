"""
Pure business logic for the PM PO Projection Mailer app.

Deliberately has no dependency on streamlit, so it can be unit-tested
and reused (e.g. from a CLI or notebook) independently of the UI.
"""

import html
import io
import json
import re
from collections import defaultdict
from datetime import date

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Business-rule constants
# ---------------------------------------------------------------------------

TIERS = {
    "CFC": [250, 500, 1000, 3000, 5000, 10000],
    "CTN": [5000, 10000, 25000, 50000, 75000, 100000, 150000, 200000],
    "TRAY": [250, 500, 1000, 3000, 5000, 10000],
    "T-SHIRT": [1000],
    "AL-WIRE": [1000],
    "ALUMINIUM WIRE": [1000],
    "BOPP": [500],
    "BOPP FILM": [500],
    "ANGLE": [500],
    "SLIP SHEET": [250],
    "SLIP-SHEET": [250],
    "THREAD": [500],
}

# Supplier roster per Item Type. TRAY intentionally shares the CFC list.
SUPPLIERS_BY_TYPE = {
    "CFC": [
        "Abirami Packaging",
        "Arr Print Pack",
        "Mayura Packaging",
        "SRI RANGA INDUSTRIES",
        "Sri Venkateswara Packings",
    ],
    "TRAY": [
        "Abirami Packaging",
        "Arr Print Pack",
        "Mayura Packaging",
        "SRI RANGA INDUSTRIES",
        "Sri Venkateswara Packings",
    ],
    "CTN": [
        "Lovely Offset Printers (P) Ltd",
        "Salem Print Pack",
        "Shraddha Saburi Printers Private Limited",
    ],
}

ITEM_TYPES = ["CFC", "CTN", "TRAY"]

# PO Issued covers deliveries due within this many days of the as-of date,
# plus anything already overdue.
PO_WINDOW_DAYS = 14

# Known corrections for items whose live description doesn't exactly match
# the master (e.g. "SC ENV" vs "DC ENV" naming variants). Extend as needed.
# Keys can be a Nav Item Code or an item name exactly as in the Requirement
# file (case/spacing ignored). Items listed in the app's "not in the Item
# Category Master" warning are the ones that need an entry here.
CATEGORY_OVERRIDES: dict = {
    # GV 25 SC ENV flavours missing from the master (only their VINDEMIA SF
    # variants are listed, under another category); confirmed by the user to
    # belong with "CTN GV MANGO 25 SC ENV" under 25TB DC ENV (AFRICA).
    "CTN GV GARDEN STRAWBERRY 25 SC ENV TBGS": "25TB DC ENV (AFRICA)",
    "CTN GV MINT 25 SC ENV TBGS": "25TB DC ENV (AFRICA)",
    "CTN GV ORANGE 25 SC ENV TBGS": "25TB DC ENV (AFRICA)",
    "CTN GV ROYAL LEMON 25 SC ENV TBGS": "25TB DC ENV (AFRICA)",
    "CTN GV SPICY GINGER 25 SC ENV TBGS": "25TB DC ENV (AFRICA)",
}


# ---------------------------------------------------------------------------
# Parsing helpers
# ---------------------------------------------------------------------------

def normalize_header(text) -> str:
    return re.sub(r"\s+", " ", (text or "").__str__().strip().upper())


def normalize_name(text) -> str:
    if text is None:
        return ""
    s = str(text).replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip().upper()


def find_header_row(ws, required_header_text: str, max_scan_rows: int = 12):
    """Scan the first rows of a worksheet for a row containing the given
    header text; return (row_number, {normalized_header: col_number})."""
    target = normalize_header(required_header_text)
    for r in range(1, min(max_scan_rows, ws.max_row) + 1):
        row_cells = list(ws[r])
        found = any(normalize_header(c.value) == target for c in row_cells)
        if found:
            col_map = {}
            for c in row_cells:
                if c.value is not None:
                    col_map[normalize_header(c.value)] = c.column
            return r, col_map
    return None, None


def col(col_map, name):
    return col_map.get(normalize_header(name))


def load_category_master(file) -> tuple:
    """Load the Item Category Master workbook (ITEM NAME, Item Type,
    ITEM CATEGORY columns, in whichever order they appear)."""
    wb = openpyxl.load_workbook(file, data_only=True)
    ws = wb["Item Category"] if "Item Category" in wb.sheetnames else wb.worksheets[0]
    header_row, col_map = find_header_row(ws, "ITEM NAME", max_scan_rows=3)
    if header_row is None:
        header_row, col_map = 1, {
            normalize_header(c.value): c.column for c in ws[1] if c.value
        }
    c_name = col(col_map, "ITEM NAME")
    c_type = col(col_map, "Item Type")
    c_cat = col(col_map, "ITEM CATEGORY")

    name_to_cat, name_to_type = {}, {}
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(row=r, column=c_name).value if c_name else None
        if not name:
            continue
        norm = normalize_name(name)
        if c_cat:
            cat = ws.cell(row=r, column=c_cat).value
            if cat:
                name_to_cat[norm] = str(cat).strip()
        if c_type:
            itype = ws.cell(row=r, column=c_type).value
            if itype:
                name_to_type[norm] = str(itype).strip().upper()
    return name_to_cat, name_to_type


def parse_requirement_workbook(file):
    """Parse the PM Requirement workbook's Shortfall + Laminates sheets.

    Returns a dict with: item_names, categories (coarse), week_shortfalls
    (per item, all 8 weeks in the current window), all_weeks (list[int]),
    projection_weeks (last 6 of the 8), laminated_codes.
    """
    wb = openpyxl.load_workbook(file, data_only=True)

    short_ws = None
    short_header = None
    for ws in wb.worksheets:
        if normalize_header(ws.title) == "SHORTFALL":
            r, cm = find_header_row(ws, "Shortfall Quantity")
            if r:
                short_ws, short_header = ws, (r, cm)
                break
    if short_ws is None:
        for ws in wb.worksheets:
            r, cm = find_header_row(ws, "Shortfall Quantity")
            if r:
                short_ws, short_header = ws, (r, cm)
                break
    if short_ws is None:
        raise ValueError("Could not find a 'Shortfall Quantity' column in the Requirement file.")

    header_row, col_map = short_header
    c_code = col(col_map, "Nav Item Code")
    c_name = col(col_map, "Item Name")
    c_cat = col(col_map, "Item Category")
    c_type = col(col_map, "Item Type")

    # Discover which weeks appear under the "Shortfall Weeks" group, in order.
    week_cols = []
    for cnum in range(1, short_ws.max_column + 1):
        header_val = short_ws.cell(row=header_row, column=cnum).value
        h = normalize_header(header_val)
        m = re.match(r"WEEK (\d+)", h)
        if m and cnum > (c_cat or 0):
            week_cols.append((int(m.group(1)), cnum))
        if len(week_cols) >= 8:
            break
    week_cols = week_cols[:8]
    all_weeks = [w for w, _ in week_cols]
    week_colmap = dict(week_cols)

    item_names, categories, item_types = {}, {}, {}
    week_shortfalls = {}
    for r in range(header_row + 1, short_ws.max_row + 1):
        code = short_ws.cell(row=r, column=c_code).value if c_code else None
        if not code or str(code).strip() in ("", "-"):
            continue
        code = str(code).strip()
        if c_name:
            item_names[code] = short_ws.cell(row=r, column=c_name).value
        if c_cat:
            categories[code] = short_ws.cell(row=r, column=c_cat).value
        if c_type:
            item_types[code] = short_ws.cell(row=r, column=c_type).value
        wk = {}
        for w, cnum in week_colmap.items():
            v = short_ws.cell(row=r, column=cnum).value
            wk[w] = v if v is not None else 0
        week_shortfalls[code] = wk

    laminated_codes = set()
    for ws in wb.worksheets:
        if "LAMIN" not in normalize_header(ws.title):
            continue
        r, cm = find_header_row(ws, "Nav Item Code")
        if not r:
            continue
        c_code2 = col(cm, "Nav Item Code")
        for rr in range(r + 1, ws.max_row + 1):
            v = ws.cell(row=rr, column=c_code2).value
            if v:
                laminated_codes.add(str(v).strip())
        break

    projection_weeks = all_weeks[2:] if len(all_weeks) >= 3 else all_weeks

    return {
        "item_names": item_names,
        "categories": categories,
        "week_shortfalls": week_shortfalls,
        "all_weeks": all_weeks,
        "projection_weeks": projection_weeks,
        "laminated_codes": laminated_codes,
    }


def parse_pending_po(file) -> pd.DataFrame:
    """Load the Pending PO export (Sheet2, or whichever sheet has
    'Outstanding Quantity'), dedupe exact-duplicate rows, and normalize
    the columns this app needs."""
    wb = openpyxl.load_workbook(file, data_only=True)
    target_ws = None
    header_row, col_map = None, None
    for ws in wb.worksheets:
        r, cm = find_header_row(ws, "Outstanding Quantity")
        if r:
            target_ws, header_row, col_map = ws, r, cm
            break
    if target_ws is None:
        raise ValueError("Could not find an 'Outstanding Quantity' column in the Pending PO file.")

    c_doc = col(col_map, "Document No")
    c_odate = col(col_map, "Order Date")
    c_cust = col(col_map, "Customer Name")
    c_item = col(col_map, "Item No")
    c_desc = col(col_map, "Item Description")
    c_deliv = col(col_map, "Delivery Date")
    c_qty = col(col_map, "Quantity")
    c_outstanding = col(col_map, "Outstanding Quantity")

    rows = []
    for r in range(header_row + 1, target_ws.max_row + 1):
        item_no = target_ws.cell(row=r, column=c_item).value if c_item else None
        if not item_no:
            continue
        rows.append(
            {
                "Document No": target_ws.cell(row=r, column=c_doc).value,
                "Order Date": target_ws.cell(row=r, column=c_odate).value,
                "Customer Name": target_ws.cell(row=r, column=c_cust).value,
                "Item No": str(item_no).strip(),
                "Item Description": target_ws.cell(row=r, column=c_desc).value,
                "Delivery Date": target_ws.cell(row=r, column=c_deliv).value,
                "Quantity": target_ws.cell(row=r, column=c_qty).value,
                "Outstanding Quantity": target_ws.cell(row=r, column=c_outstanding).value,
            }
        )
    df = pd.DataFrame(rows).drop_duplicates().reset_index(drop=True)
    df["Order Date"] = pd.to_datetime(df["Order Date"], dayfirst=True, errors="coerce")
    df["Delivery Date"] = pd.to_datetime(df["Delivery Date"], dayfirst=True, errors="coerce")
    return df


# ---------------------------------------------------------------------------
# Business logic
# ---------------------------------------------------------------------------

def compute_moq(item_code, category, qty, item_names, laminated_codes):
    if qty is None or qty <= 0:
        return None
    cat_norm = (category or "").strip().upper()
    name_up = (item_names.get(item_code) or "").upper()
    if item_code in laminated_codes:
        tiers = [300]
    elif cat_norm == "POUCH" and "GARANT" in name_up:
        tiers = [5000, 10000, 25000, 50000]
    else:
        tiers = TIERS.get(cat_norm)
    if not tiers:
        return None
    for t in tiers:
        if qty <= t:
            return t
    return round(qty * 1.02)


def _cat_type(category, master_name_to_cat, master_name_to_type):
    """Item Type used by the master for a given category (first match)."""
    for name, cat in master_name_to_cat.items():
        if cat == category and master_name_to_type.get(name):
            return master_name_to_type[name]
    return None


_OVERRIDES_BY_NAME = {normalize_name(k): v for k, v in CATEGORY_OVERRIDES.items()}


def build_item_category_map(req_data, master_name_to_cat, master_name_to_type):
    """Map each Nav Item Code to (fine category, item type)."""
    item_names = req_data["item_names"]
    categories = req_data["categories"]
    item_to_cat, item_to_type = {}, {}
    for code, name in item_names.items():
        norm = normalize_name(name)
        fine_cat = master_name_to_cat.get(norm)
        override = CATEGORY_OVERRIDES.get(code) or _OVERRIDES_BY_NAME.get(norm)
        if override:
            fine_cat = override
        if fine_cat:
            item_to_cat[code] = fine_cat
            item_to_type[code] = master_name_to_type.get(norm) or _cat_type(fine_cat, master_name_to_cat, master_name_to_type) \
                or (categories.get(code) or "").strip().upper()
    return item_to_cat, item_to_type


def build_item_week_projections(req_data):
    item_names = req_data["item_names"]
    categories = req_data["categories"]
    laminated_codes = req_data["laminated_codes"]
    week_shortfalls = req_data["week_shortfalls"]
    projection_weeks = req_data["projection_weeks"]

    item_week_proj = {}
    for code, wk in week_shortfalls.items():
        cat = categories.get(code)
        item_week_proj[code] = {
            w: compute_moq(code, cat, wk.get(w, 0), item_names, laminated_codes)
            for w in projection_weeks
        }
    return item_week_proj


def build_unmatched_items_df(req_data, item_week_proj, item_to_cat) -> pd.DataFrame:
    """CFC/CTN/TRAY items that have a projection but no Item Category (not in
    the master and no override), so they are missing from PO Projection."""
    weeks = req_data["projection_weeks"]
    rows = []
    for code, projs in item_week_proj.items():
        coarse = (req_data["categories"].get(code) or "").strip().upper()
        total = sum(v or 0 for v in projs.values())
        if code in item_to_cat or coarse not in ITEM_TYPES or total <= 0:
            continue
        row = {"Nav Item Code": code, "Item Name": req_data["item_names"].get(code), "Item Type": coarse}
        row.update({f"Wk #{w}": projs.get(w) or 0 for w in weeks})
        row["Total"] = total
        rows.append(row)
    cols = ["Nav Item Code", "Item Name", "Item Type"] + [f"Wk #{w}" for w in weeks] + ["Total"]
    return pd.DataFrame(rows, columns=cols).sort_values("Total", ascending=False).reset_index(drop=True)


def build_po_issued_df(po_df: pd.DataFrame, today: date, window_days: int) -> pd.DataFrame:
    today_ts = pd.Timestamp(today)
    cutoff = today_ts + pd.Timedelta(days=window_days)
    window = po_df[po_df["Delivery Date"] <= cutoff].copy()
    window["No. of days to arrive"] = (window["Delivery Date"] - today_ts).dt.days
    window = window.sort_values(["Customer Name", "Item No", "Delivery Date"]).reset_index(drop=True)
    return window


def build_category_table(item_week_proj, item_to_cat, item_to_type, category_order, projection_weeks):
    """One row per category (in category_order), Item Type, week columns, Total."""
    cat_week_sum = defaultdict(lambda: {w: 0.0 for w in projection_weeks})
    cat_type = {}
    for code, projs in item_week_proj.items():
        cat = item_to_cat.get(code)
        if not cat:
            continue
        cat_type[cat] = item_to_type.get(code, cat_type.get(cat, ""))
        for w in projection_weeks:
            v = projs.get(w)
            if v:
                cat_week_sum[cat][w] += v

    rows = []
    for cat in category_order:
        weeksum = cat_week_sum.get(cat, {w: 0.0 for w in projection_weeks})
        total = sum(weeksum.values())
        row = {"Item Category": cat, "Item Type": cat_type.get(cat, "")}
        for w in projection_weeks:
            row[f"Wk #{w}"] = weeksum[w]
        row["Total"] = total
        rows.append(row)
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Excel report generation
# ---------------------------------------------------------------------------

def style_header_cell(cell, fill_hex="F4B183"):
    thin = Side(style="thin", color="000000")
    cell.font = Font(name="Calibri", bold=True, size=11)
    cell.fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_data_cell(cell, bold=False, color=None, fill_hex=None):
    thin = Side(style="thin", color="000000")
    cell.font = Font(name="Calibri", size=11, bold=bold, color=color)
    if fill_hex:
        cell.fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def write_supplier_projection_sheet(ws, alloc_df: pd.DataFrame, projection_weeks) -> None:
    """Same layout as the in-app supplier-wise projection: supplier name merged
    down its rows, Item category, Item type, one column per week."""
    headers = ["Supplier", "Item category", "Item type"] + [week_label(w) for w in projection_weeks]
    ws.append(headers)
    for cidx in range(1, len(headers) + 1):
        style_header_cell(ws.cell(row=1, column=cidx))

    r = 2
    for supplier, grp in alloc_df.groupby("Supplier", sort=False):
        first = r
        for _, row in grp.iterrows():
            style_data_cell(ws.cell(row=r, column=1, value=supplier))
            c = ws.cell(row=r, column=2, value=row["Item Category"])
            style_data_cell(c)
            c.alignment = Alignment(horizontal="left", vertical="center")
            style_data_cell(ws.cell(row=r, column=3, value=row["Item Type"]))
            for wi, w in enumerate(projection_weeks):
                v = row[f"Wk #{w}"]
                c = ws.cell(row=r, column=4 + wi, value=float(v) if v else None)
                style_data_cell(c)
                c.number_format = "#,##0"
                c.alignment = Alignment(horizontal="right", vertical="center")
            r += 1
        if r - 1 > first:
            ws.merge_cells(start_row=first, start_column=1, end_row=r - 1, end_column=1)

    if alloc_df.empty:
        ws.cell(row=2, column=1, value="No suppliers assigned yet.").font = Font(name="Calibri", size=11, italic=True)

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 11
    for i in range(len(projection_weeks)):
        ws.column_dimensions[get_column_letter(4 + i)].width = 11
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "B2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def generate_report_workbook(po_issued: pd.DataFrame, category_df: pd.DataFrame, projection_weeks,
                             alloc_df: pd.DataFrame = None) -> bytes:
    """PO Issued + PO Projection sheets; with `alloc_df` (from
    allocate_to_suppliers) also a Supplier-wise Projection sheet."""
    wb = openpyxl.Workbook()

    # ---- PO Issued ----
    ws1 = wb.active
    ws1.title = "PO Issued"
    headers1 = ["Supplier Name", "PO number", "PO date", "Item Description",
                "PO Qty", "Outstanding Qty", "Delivery date", "No. of days to arrive"]
    ws1.append(headers1)
    for cidx in range(1, len(headers1) + 1):
        style_header_cell(ws1.cell(row=1, column=cidx))

    for i, row in po_issued.iterrows():
        r = i + 2
        ws1.cell(row=r, column=1, value=row["Customer Name"])
        style_data_cell(ws1.cell(row=r, column=1))
        ws1.cell(row=r, column=1).alignment = Alignment(vertical="center")
        ws1.cell(row=r, column=2, value=row["Document No"])
        style_data_cell(ws1.cell(row=r, column=2))
        c = ws1.cell(row=r, column=3, value=row["Order Date"])
        style_data_cell(c)
        c.number_format = "DD/MM/YYYY"
        ws1.cell(row=r, column=4, value=row["Item Description"])
        style_data_cell(ws1.cell(row=r, column=4))
        ws1.cell(row=r, column=4).alignment = Alignment(vertical="center")
        c = ws1.cell(row=r, column=5, value=float(row["Quantity"]) if pd.notna(row["Quantity"]) else None)
        style_data_cell(c)
        c.number_format = "#,##0"
        c = ws1.cell(row=r, column=6, value=float(row["Outstanding Quantity"]) if pd.notna(row["Outstanding Quantity"]) else None)
        style_data_cell(c)
        c.number_format = "#,##0"
        c = ws1.cell(row=r, column=7, value=row["Delivery Date"])
        style_data_cell(c)
        c.number_format = "DD/MM/YYYY"
        days = int(row["No. of days to arrive"])
        delayed = days < 0
        c = ws1.cell(row=r, column=8, value=days)
        style_data_cell(c, bold=delayed, color="FFFF0000" if delayed else None,
                         fill_hex="FCE5E5" if delayed else None)

    widths1 = {"A": 30, "B": 20, "C": 12, "D": 42, "E": 10, "F": 12, "G": 13, "H": 12}
    for c, w in widths1.items():
        ws1.column_dimensions[c].width = w
    ws1.row_dimensions[1].height = 30
    last1 = ws1.max_row
    ws1.auto_filter.ref = f"A1:H{last1}"
    ws1.freeze_panes = "A2"
    ws1.page_setup.orientation = "landscape"
    ws1.page_setup.fitToWidth = 1
    ws1.page_setup.fitToHeight = 0
    ws1.sheet_properties.pageSetUpPr.fitToPage = True

    # ---- PO Projection (only categories with a projection) ----
    ws2 = wb.create_sheet("PO Projection")
    nonzero_df = category_df[category_df["Total"] > 0].reset_index(drop=True)

    headers2 = ["Item Category", "Item Type"] + [f"Wk #{w}" for w in projection_weeks] + ["Total"]
    ws2.append(headers2)
    for cidx in range(1, len(headers2) + 1):
        style_header_cell(ws2.cell(row=1, column=cidx))

    for i, row in nonzero_df.iterrows():
        r = i + 2
        ws2.cell(row=r, column=1, value=row["Item Category"])
        style_data_cell(ws2.cell(row=r, column=1))
        ws2.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws2.cell(row=r, column=2, value=row["Item Type"])
        style_data_cell(ws2.cell(row=r, column=2))
        for wi, w in enumerate(projection_weeks):
            val = row[f"Wk #{w}"]
            c = ws2.cell(row=r, column=3 + wi)
            if val and val > 0:
                c.value = float(val)
                c.number_format = "#,##0"
            else:
                c.value = "-"
            style_data_cell(c)
        c = ws2.cell(row=r, column=3 + len(projection_weeks), value=float(row["Total"]))
        c.number_format = "#,##0"
        style_data_cell(c, bold=True)

    widths2 = {"A": 32, "B": 12}
    for c, w in widths2.items():
        ws2.column_dimensions[c].width = w
    for i in range(len(projection_weeks) + 1):
        ws2.column_dimensions[get_column_letter(3 + i)].width = 13
    ws2.row_dimensions[1].height = 24
    last2 = ws2.max_row
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers2))}{last2}"
    ws2.freeze_panes = "A2"
    ws2.page_setup.orientation = "landscape"
    ws2.page_setup.fitToWidth = 1
    ws2.page_setup.fitToHeight = 0
    ws2.sheet_properties.pageSetUpPr.fitToPage = True

    if alloc_df is not None:
        write_supplier_projection_sheet(wb.create_sheet("Supplier-wise Projection"), alloc_df, projection_weeks)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Email
# ---------------------------------------------------------------------------

EMAIL_TEMPLATE_INTRO = (
    "Please find the Item category wise projections for upcoming weeks.\n\n"
    "Note: Projection Qty is only to Procure raw materials. Upon the receipt "
    "of PO, you can print as per the delivery schedule."
)


# ---------------------------------------------------------------------------
# Supplier allocation (per category/week cell)
# ---------------------------------------------------------------------------

def allocate_to_suppliers(category_df: pd.DataFrame, selections: dict, projection_weeks) -> pd.DataFrame:
    """Split each category/week projection evenly across the suppliers picked
    for that cell. `selections` maps (item_category, week) -> [supplier, ...].

    Returns one row per (Supplier, Item Category) with the allocated quantity
    in each "Wk #N" column, ordered by supplier, then Item Type, then category.
    """
    week_cols = [f"Wk #{w}" for w in projection_weeks]
    by_cat = category_df.set_index("Item Category")
    alloc = defaultdict(lambda: {c: 0.0 for c in week_cols})
    cat_types = {}
    for (cat, w), suppliers in selections.items():
        suppliers = [s for s in dict.fromkeys(suppliers or []) if s]
        if not suppliers or cat not in by_cat.index:
            continue
        qty = by_cat.at[cat, f"Wk #{w}"]
        if not qty or qty <= 0:
            continue
        share = qty / len(suppliers)
        for s in suppliers:
            alloc[(s, cat)][f"Wk #{w}"] += share
            cat_types[cat] = by_cat.at[cat, "Item Type"]

    columns = ["Supplier", "Item Category", "Item Type"] + week_cols
    if not alloc:
        return pd.DataFrame(columns=columns)

    supplier_order = list(dict.fromkeys(s for sups in SUPPLIERS_BY_TYPE.values() for s in sups))
    type_order = {t: i for i, t in enumerate(ITEM_TYPES)}

    def sort_key(key):
        s, cat = key
        s_idx = supplier_order.index(s) if s in supplier_order else len(supplier_order)
        return (s_idx, s, type_order.get(cat_types[cat], len(type_order)), cat)

    rows = []
    for key in sorted(alloc, key=sort_key):
        s, cat = key
        row = {"Supplier": s, "Item Category": cat, "Item Type": cat_types[cat]}
        row.update({c: round(v) for c, v in alloc[key].items()})
        rows.append(row)
    return pd.DataFrame(rows, columns=columns)


def week_label(w) -> str:
    return f"Wk {w}"


def build_allocation_html(alloc_df: pd.DataFrame, projection_weeks) -> str:
    """Render the supplier allocation as a bordered HTML table: supplier name
    merged down its rows, then Item category, Item type, and one column per
    week (blank where nothing is allocated). Inline styles so it survives
    email clients."""
    cell = "border:1px solid #000;padding:4px 8px;font-family:Calibri,Arial,sans-serif;font-size:14px;"
    head = ["Item category", "Item type"] + [week_label(w) for w in projection_weeks]
    out = ['<table style="border-collapse:collapse;">', "<tr>", f'<th style="{cell}"></th>']
    out += [f'<th style="{cell}text-align:left;font-weight:normal;">{h}</th>' for h in head]
    out.append("</tr>")
    for supplier, grp in alloc_df.groupby("Supplier", sort=False):
        for i, (_, row) in enumerate(grp.iterrows()):
            out.append("<tr>")
            if i == 0:
                out.append(
                    f'<td rowspan="{len(grp)}" style="{cell}text-align:center;vertical-align:middle;">'
                    f"{html.escape(supplier)}</td>"
                )
            out.append(f'<td style="{cell}">{html.escape(str(row["Item Category"]))}</td>')
            out.append(f'<td style="{cell}text-align:center;">{html.escape(str(row["Item Type"]))}</td>')
            for w in projection_weeks:
                v = row[f"Wk #{w}"]
                out.append(f'<td style="{cell}text-align:right;">{f"{v:.0f}" if v else ""}</td>')
            out.append("</tr>")
    out.append("</table>")
    return "".join(out)


_COMPANY_SUFFIXES = r"\b(PRIVATE|PVT|LIMITED|LTD|P|CO|COMPANY|INDIA)\b"


def _supplier_key(name) -> str:
    """Loose key for matching supplier names across files: case, punctuation
    and company suffixes (Pvt Ltd, (P) Ltd, Private Limited...) ignored."""
    s = re.sub(r"[^A-Z0-9 ]", " ", normalize_name(name))
    s = re.sub(_COMPANY_SUFFIXES, " ", s)
    return re.sub(r"\s+", " ", s).strip()


def match_roster_supplier(name):
    """Map a supplier name from the Pending PO file to its name in
    SUPPLIERS_BY_TYPE (exact loose match first, then prefix match either way).
    Returns None when the supplier isn't on the roster."""
    key = _supplier_key(name)
    if not key:
        return None
    roster = list(dict.fromkeys(s for sups in SUPPLIERS_BY_TYPE.values() for s in sups))
    keys = {s: _supplier_key(s) for s in roster}
    for s, k in keys.items():
        if k == key:
            return s
    for s, k in keys.items():
        shorter = min(k, key, key=len)
        # Prefix match only on whole words, and only for 2+ word names, so a
        # bare "SRI" can't grab "SRI RANGA INDUSTRIES".
        if len(shorter.split()) >= 2 and (key.startswith(k + " ") or k.startswith(key + " ")):
            return s
    return None


PO_EMAIL_INTRO = "Please find the below list of pending purchase orders as on today ({date})."
EMAIL_SIGN_OFF = ["Thank you.", "MJIL - Packing Materials"]


def _fmt_date(v) -> str:
    return v.strftime("%d/%m/%Y") if pd.notna(v) else ""


def _fmt_qty(v) -> str:
    return f"{float(v):,.0f}" if pd.notna(v) else ""


def build_po_issued_html(po_rows: pd.DataFrame) -> str:
    """Supplier's pending POs as a bordered table; negative "No. of days to
    arrive" in bold red on a red fill, matching the PO Issued sheet."""
    cell = "border:1px solid #000;padding:4px 8px;font-family:Calibri,Arial,sans-serif;font-size:14px;"
    head = ["PO number", "PO date", "Item Description", "PO Qty", "Outstanding Qty",
            "Delivery date", "No. of days to arrive"]
    out = ['<table style="border-collapse:collapse;">', "<tr>"]
    out += [f'<th style="{cell}text-align:left;font-weight:normal;">{h}</th>' for h in head]
    out.append("</tr>")
    for _, row in po_rows.iterrows():
        days = row["No. of days to arrive"]
        late = pd.notna(days) and days < 0
        days_style = "color:#FF0000;font-weight:bold;background:#FCE5E5;" if late else ""
        out.append("<tr>")
        out.append(f'<td style="{cell}">{html.escape(str(row["Document No"] or ""))}</td>')
        out.append(f'<td style="{cell}text-align:center;">{_fmt_date(row["Order Date"])}</td>')
        out.append(f'<td style="{cell}">{html.escape(str(row["Item Description"] or ""))}</td>')
        out.append(f'<td style="{cell}text-align:right;">{_fmt_qty(row["Quantity"])}</td>')
        out.append(f'<td style="{cell}text-align:right;">{_fmt_qty(row["Outstanding Quantity"])}</td>')
        out.append(f'<td style="{cell}text-align:center;">{_fmt_date(row["Delivery Date"])}</td>')
        out.append(f'<td style="{cell}text-align:center;{days_style}">{"" if pd.isna(days) else int(days)}</td>')
        out.append("</tr>")
    out.append("</table>")
    return "".join(out)


def _po_intro(as_of) -> str:
    return PO_EMAIL_INTRO.format(date=pd.Timestamp(as_of or date.today()).strftime("%d/%m/%Y"))


def build_supplier_email_html(supplier_name: str, supplier_alloc_df: pd.DataFrame, projection_weeks,
                              supplier_po_df: pd.DataFrame = None, as_of=None) -> str:
    parts = [
        '<div style="font-family:Calibri,Arial,sans-serif;font-size:14px;">',
        f"<p>Hello {html.escape(supplier_name)},</p>",
    ]
    if supplier_po_df is not None and not supplier_po_df.empty:
        parts += [f"<p>{html.escape(_po_intro(as_of))}</p>",
                  "<p><b>Pending deliveries</b></p>",
                  build_po_issued_html(supplier_po_df)]
    if not supplier_alloc_df.empty:
        intro, _, note = EMAIL_TEMPLATE_INTRO.partition("\n\n")
        parts += [f"<p><b>PO Projection</b><br>{html.escape(intro)}</p>",
                  build_allocation_html(supplier_alloc_df, projection_weeks),
                  f'<p style="color:#FF0000;font-weight:bold;">{html.escape(note)}</p>']
    parts.append("<p>" + "<br>".join(html.escape(l) for l in EMAIL_SIGN_OFF) + "</p></div>")
    return "".join(parts)


def build_supplier_allocation_text(supplier_name: str, supplier_alloc_df: pd.DataFrame, projection_weeks,
                                   supplier_po_df: pd.DataFrame = None, as_of=None) -> str:
    """Plain-text fallback of the HTML email, for clients that don't render HTML."""
    lines = [f"Hello {supplier_name},", ""]
    if supplier_po_df is not None and not supplier_po_df.empty:
        lines += [_po_intro(as_of), "", "Pending deliveries"]
        for _, row in supplier_po_df.iterrows():
            days = row["No. of days to arrive"]
            lines.append(
                f"- PO {row['Document No']} ({_fmt_date(row['Order Date'])}): {row['Item Description']} | "
                f"PO Qty {_fmt_qty(row['Quantity'])} | Outstanding {_fmt_qty(row['Outstanding Quantity'])} | "
                f"Delivery {_fmt_date(row['Delivery Date'])} | Days to arrive {'' if pd.isna(days) else int(days)}"
            )
        lines.append("")
    if not supplier_alloc_df.empty:
        intro, _, note = EMAIL_TEMPLATE_INTRO.partition("\n\n")
        lines += ["PO Projection", intro, ""]
        for _, row in supplier_alloc_df.iterrows():
            lines.append(f"- {row['Item Category']} ({row['Item Type']})")
            for w in projection_weeks:
                v = row[f"Wk #{w}"]
                if v:
                    lines.append(f"    {week_label(w)}: {v:.0f}")
        lines += ["", note, ""]
    lines += EMAIL_SIGN_OFF
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Supplier email IDs (JSON file store)
# ---------------------------------------------------------------------------

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def all_suppliers() -> list:
    return list(dict.fromkeys(s for sups in SUPPLIERS_BY_TYPE.values() for s in sups))


def is_valid_email(addr: str) -> bool:
    return bool(EMAIL_RE.match((addr or "").strip()))


def load_supplier_emails(path, seed=None) -> dict:
    """Return {supplier: [email, ...]} from the JSON file at `path`. If the
    file doesn't exist yet, start from `seed` ({supplier: "a@x, b@y" or list}),
    e.g. the supplier_emails table in Streamlit secrets."""
    try:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        data = {}
        for supplier, value in (seed or {}).items():
            items = value.split(",") if isinstance(value, str) else list(value)
            data[supplier] = [e.strip() for e in items if is_valid_email(e)]
    return {s: list(dict.fromkeys(data.get(s, []))) for s in set(all_suppliers()) | set(data)}


def save_supplier_emails(path, emails: dict) -> None:
    with open(path, "w", encoding="utf-8") as f:
        json.dump({s: v for s, v in emails.items() if v}, f, indent=2, sort_keys=True)
